import os
import io
import re
import pandas as pd
from typing import List, Dict, Any, Tuple
try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import docx
except ImportError:
    docx = None

class DocumentParser:
    """
    Document Ingestion and Parsing Module.
    Supports PDF, DOCX, CSV, Plain Text, and Markdown.
    Extracts text content along with rich structural metadata.
    """
    
    @staticmethod
    def parse_file(file_bytes: bytes, filename: str) -> Dict[str, Any]:
        ext = os.path.splitext(filename)[1].lower()
        
        if ext == ".pdf":
            text, pages_meta = DocumentParser._parse_pdf(file_bytes)
            doc_type = "PDF Document"
        elif ext == ".docx":
            text, pages_meta = DocumentParser._parse_docx(file_bytes)
            doc_type = "Word Document (DOCX)"
        elif ext == ".csv":
            text, pages_meta = DocumentParser._parse_csv(file_bytes)
            doc_type = "CSV Spreadsheet"
        elif ext in [".txt", ".md", ".json", ".yaml"]:
            text, pages_meta = DocumentParser._parse_text(file_bytes)
            doc_type = "Plain Text / Markdown"
        else:
            # Fallback to UTF-8 text decoding
            text, pages_meta = DocumentParser._parse_text(file_bytes)
            doc_type = "Raw Text"
            
        word_count = len(re.findall(r'\w+', text))
        line_count = len(text.splitlines())
        
        return {
            "filename": filename,
            "doc_type": doc_type,
            "extension": ext,
            "raw_text": text,
            "word_count": word_count,
            "line_count": line_count,
            "pages_meta": pages_meta,
            "file_size_bytes": len(file_bytes)
        }

    @staticmethod
    def _parse_pdf(file_bytes: bytes) -> Tuple[str, List[Dict[str, Any]]]:
        if not pypdf:
            # Fallback text extraction if pypdf is not available
            text = file_bytes.decode('utf-8', errors='ignore')
            return text, [{"page": 1, "text": text}]
            
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        full_text_parts = []
        pages_meta = []
        
        for idx, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            page_num = idx + 1
            full_text_parts.append(f"--- Page {page_num} ---\n{page_text}")
            pages_meta.append({"page": page_num, "text": page_text})
            
        return "\n\n".join(full_text_parts), pages_meta

    @staticmethod
    def _parse_docx(file_bytes: bytes) -> Tuple[str, List[Dict[str, Any]]]:
        if not docx:
            text = file_bytes.decode('utf-8', errors='ignore')
            return text, [{"page": 1, "text": text}]
            
        doc = docx.Document(io.BytesIO(file_bytes))
        text_parts = []
        
        for p in doc.paragraphs:
            if p.text.strip():
                text_parts.append(p.text.strip())
                
        for table in doc.tables:
            table_text = []
            for row in table.rows:
                row_str = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_str:
                    table_text.append(row_str)
            if table_text:
                text_parts.append("[TABLE]\n" + "\n".join(table_text))
                
        full_text = "\n\n".join(text_parts)
        return full_text, [{"section": "Document Body", "text": full_text}]

    @staticmethod
    def _parse_csv(file_bytes: bytes) -> Tuple[str, List[Dict[str, Any]]]:
        try:
            df = pd.read_csv(io.BytesIO(file_bytes))
            formatted_lines = []
            formatted_lines.append(f"CSV Headers: {', '.join(df.columns.astype(str))}")
            formatted_lines.append(f"Total Rows: {len(df)}")
            formatted_lines.append("Data Summary:")
            
            for idx, row in df.iterrows():
                row_items = [f"{col}: {val}" for col, val in row.items() if pd.notna(val)]
                formatted_lines.append(f"Row {idx+1}: " + " | ".join(row_items))
                
            full_text = "\n".join(formatted_lines)
            return full_text, [{"sheet": "Main", "text": full_text}]
        except Exception as e:
            raw = file_bytes.decode('utf-8', errors='ignore')
            return raw, [{"sheet": "Raw CSV", "text": raw}]

    @staticmethod
    def _parse_text(file_bytes: bytes) -> Tuple[str, List[Dict[str, Any]]]:
        text = file_bytes.decode('utf-8', errors='ignore')
        return text, [{"section": "Full Text", "text": text}]
